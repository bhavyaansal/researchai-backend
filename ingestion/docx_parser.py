"""
DOCX parsing using python-docx.

Walks the document body in true reading order (paragraphs, images, and
tables interleaved as they actually appear), not just paragraphs.text
joined together. Extracts:
  - paragraph text (as before)
  - inline images -> saved to <assets_dir>/figures/, replaced in the
    text stream with a placeholder token: [[FIGURE:fig_0001.png]]
  - tables -> saved as JSON (rows of cell text) to
    <assets_dir>/tables/, replaced with a placeholder token:
    [[TABLE:tbl_0001.json]]

Downstream rendering (ieee_pdf.py) resolves these placeholder tokens
back into real images/tables when rebuilding the document. Coordinate
mapping (start_char/end_char) still works fine on the resulting text —
a placeholder is just another run of characters like any other.
"""
import os
import json
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def _iter_block_items(document):
    """
    Yield each paragraph and table in the document body, in the order
    they actually appear (python-docx's own .paragraphs and .tables
    properties don't preserve relative order between the two).
    """
    parent_elm = document.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _extract_run_images(run, document, figures_dir, counter):
    """
    Finds embedded images referenced within a single run (via <a:blip>
    elements), saves each to figures_dir, and returns a list of
    placeholder tokens (one per image found in this run).
    """
    tokens = []
    blips = run._element.findall(".//" + qn("a:blip"))
    for blip in blips:
        r_id = blip.get(qn("r:embed"))
        if not r_id or r_id not in document.part.rels:
            continue
        rel = document.part.rels[r_id]
        if "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
        except Exception:
            continue

        counter["n"] += 1
        ext = os.path.splitext(rel.target_part.partname)[1] or ".png"
        fname = f"fig_{counter['n']:04d}{ext}"
        os.makedirs(figures_dir, exist_ok=True)
        with open(os.path.join(figures_dir, fname), "wb") as f:
            f.write(blob)
        tokens.append(f"[[FIGURE:{fname}]]")
    return tokens


def _paragraph_to_text(paragraph, document, figures_dir, counter):
    """
    Builds one paragraph's text, inline with any image placeholders in
    the position they actually occur among the paragraph's runs.
    """
    pieces = []
    for run in paragraph.runs:
        if run.text:
            pieces.append(run.text)
        img_tokens = _extract_run_images(run, document, figures_dir, counter)
        for token in img_tokens:
            # Give each figure its own paragraph so it renders as a
            # standalone block, not inline mid-sentence.
            pieces.append(f"\n\n{token}\n\n")
    return "".join(pieces).strip()


def _table_to_placeholder(table, tables_dir, counter):
    """
    Extracts a table's cell text into a JSON file and returns the
    placeholder token referencing it.
    """
    rows = []
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])

    counter["n"] += 1
    fname = f"tbl_{counter['n']:04d}.json"
    os.makedirs(tables_dir, exist_ok=True)
    with open(os.path.join(tables_dir, fname), "w", encoding="utf-8") as f:
        json.dump(rows, f)
    return f"[[TABLE:{fname}]]"


def parse_docx(file_path: str, assets_dir: str = None) -> str:
    """
    Extract all paragraph text, inline images, and tables from a Word
    document, in true reading order.

    If assets_dir is given, images are saved under
    <assets_dir>/figures/ and tables under <assets_dir>/tables/, with
    placeholder tokens inserted into the returned text at the correct
    position. If assets_dir is None, images/tables are skipped
    entirely and only paragraph text is returned (old behavior).
    """
    doc = Document(file_path)

    if not assets_dir:
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paragraphs)

    figures_dir = os.path.join(assets_dir, "figures")
    tables_dir = os.path.join(assets_dir, "tables")
    fig_counter = {"n": 0}
    tbl_counter = {"n": 0}

    blocks = []
    for item in _iter_block_items(doc):
        if isinstance(item, Paragraph):
            text = _paragraph_to_text(item, doc, figures_dir, fig_counter)
            if text:
                blocks.append(text)
        elif isinstance(item, Table):
            token = _table_to_placeholder(item, tables_dir, tbl_counter)
            blocks.append(token)

    return "\n\n".join(blocks)